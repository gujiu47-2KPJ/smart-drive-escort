#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成《智行护航 · 技术路线与调试记录》DOCX

设计规范（Documents skill）：
  preset      : compact_reference_guide
  header page : memo_masthead
  页面        : US Letter 纵向，1.0in 页边距，正文宽 6.5in / 9360 DXA
  字体        : 拉丁 Calibri 11pt；中文 eastAsia = Microsoft YaHei
  正文        : 左对齐，段前 0，段后 6pt，行距 1.25
  标题        : H1 16pt #2E74B5 / H2 13pt #2E74B5 / H3 12pt #1F4D78（加粗）
  列表        : 真编号定义，marker 0.187in，text 0.375in，hanging 0.188in，段后 4pt，行距 1.25
  表格        : 宽 9360 DXA，tblInd 120 DXA，单元格边距 80/80/120/120，表头填充 #E8EEF5
  具名覆盖    : (1) 表格正文 10pt    (2) 中文字体 Microsoft YaHei    (3) 首屏底部分隔线
"""

import os
import re
import sys

SKILL = r"C:\Users\gujiu\.codex\plugins\cache\openai-primary-runtime\documents\26.812.11052\skills\documents"
sys.path.insert(0, os.path.join(SKILL, "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor, Twips  # noqa: E402

# ----------------------------------------------------------------- tokens
LATIN = "Calibri"
CJK = "Microsoft YaHei"

BODY_SIZE = 11
BODY_AFTER = 6
BODY_LINE = 1.25

H1 = dict(size=16, color="2E74B5", before=18, after=10)
H2 = dict(size=13, color="2E74B5", before=14, after=7)
H3 = dict(size=12, color="1F4D78", before=10, after=5)

TABLE_WIDTH = 9360
TABLE_INDENT = 120
TABLE_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}
TABLE_FILL = "E8EEF5"
TABLE_TEXT_SIZE = 10          # named override: table body 10pt

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0x00, 0x00, 0x00)
GOLD = RGBColor(0x7A, 0x5A, 0x00)
RED = RGBColor(0x9B, 0x1C, 0x1C)
CALLOUT_FILL = "F4F6F9"

TITLE_TEXT = "智行护航 · 技术路线与调试记录"
SUBTITLE_TEXT = "ESP32-S3 端侧多模态融合主动安全头盔系统"
RUNNING_LABEL = "智行护航 · 技术路线与调试记录"
RUNNING_RIGHT = "会话技术总结"


# ----------------------------------------------------------------- helpers
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(text):
    """剔除 XML 不允许的控制字符（\v 这类真实转义会直接让 lxml 报错）。"""
    return _CTRL.sub("", str(text))


def set_font(run, size=None, color=None, bold=None, italic=None, latin=LATIN, cjk=CJK):
    if run.text:
        run.text = clean(run.text)
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def style_font(style, size, color=None, bold=None, latin=LATIN, cjk=CJK):
    style.font.name = latin
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)


def para(doc, text="", size=BODY_SIZE, color=None, bold=None, italic=None,
         before=None, after=BODY_AFTER, line=BODY_LINE,
         align=WD_ALIGN_PARAGRAPH.LEFT, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        set_font(p.add_run(clean(text)), size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    p.alignment = align
    return p


def rich(doc, parts, size=BODY_SIZE, before=None, after=BODY_AFTER, line=BODY_LINE):
    """parts: list of (text, bold, color)"""
    p = doc.add_paragraph()
    for text, bold, color in parts:
        set_font(p.add_run(clean(text)), size=size, bold=bold, color=color)
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return p


def shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


# OOXML 对 pPr / trPr 的子元素顺序有强制要求，顺序错了 Word 会直接忽略该元素。
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]


def insert_ordered(parent, child, order):
    """按 OOXML 规定的顺序把 child 插入 parent，而不是简单 append。"""
    tag = child.tag.split("}")[-1]
    idx = order.index(tag) if tag in order else len(order)
    for existing in parent:
        etag = existing.tag.split("}")[-1]
        eidx = order.index(etag) if etag in order else len(order)
        if eidx > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def para_shading(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, PPR_ORDER)


def para_left_bar(p, color="2E74B5", width=18):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(width))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pbdr.append(left)
    insert_ordered(pPr, pbdr, PPR_ORDER)


def para_bottom_rule(p, color="B9C6D6", width=8):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(width))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    insert_ordered(pPr, pbdr, PPR_ORDER)


def callout(doc, label, text, color=GOLD):
    p = doc.add_paragraph()
    set_font(p.add_run(clean(label)), size=10.5, bold=True, color=color)
    set_font(p.add_run(clean(text)), size=10.5, color=BLACK)
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.25
    pf.left_indent = Twips(140)
    pf.right_indent = Twips(60)
    para_shading(p, CALLOUT_FILL)
    para_left_bar(p, "7A5A00")
    return p


# ----------------------------------------------------------------- numbering
def add_numbering(doc):
    """建立真实编号定义：bullet L0 / decimal L0，并按预设缩进。"""
    numbering = doc.part.numbering_part.element

    def abstract(aid, fmt, text, font=None):
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "singleLevel")
        a.append(mlt)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        nf = OxmlElement("w:numFmt")
        nf.set(qn("w:val"), fmt)
        lvl.append(nf)
        lt = OxmlElement("w:lvlText")
        lt.set(qn("w:val"), text)
        lvl.append(lt)
        lj = OxmlElement("w:lvlJc")
        lj.set(qn("w:val"), "left")
        lvl.append(lj)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")          # 0.375in
        ind.set(qn("w:hanging"), "270")       # 0.188in
        ppr.append(ind)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            rf.set(qn("w:ascii"), font)
            rf.set(qn("w:hAnsi"), font)
            rf.set(qn("w:eastAsia"), CJK)
            rpr.append(rf)
            lvl.append(rpr)
        a.append(lvl)
        return a

    bullet_aid, number_aid = 90, 91
    # 关键：所有 w:abstractNum 必须排在所有 w:num 之前，否则 Word 忽略整段编号定义。
    first_num = None
    for child in numbering:
        if child.tag == qn("w:num"):
            first_num = child
            break
    # marker 用 Calibri 的 • ；不要用 Symbol 字体，它没有 U+2022 字形，会渲染成空心方框。
    new_abstracts = [abstract(bullet_aid, "bullet", "\u2022", font=None),
                     abstract(number_aid, "decimal", "%1.", font=LATIN)]
    if first_num is not None:
        for a in new_abstracts:
            first_num.addprevious(a)
    else:
        for a in new_abstracts:
            numbering.append(a)

    for nid, aid in ((90, bullet_aid), (91, number_aid)):
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(nid))
        ai = OxmlElement("w:abstractNumId")
        ai.set(qn("w:val"), str(aid))
        n.append(ai)
        numbering.append(n)

    return {"bullet": 90, "number": 91}


def add_list_styles(doc, num_ids):
    for name, num_id, base in (("List Tighter Bullet", num_ids["bullet"], "List Bullet"),
                               ("List Tighter Number", num_ids["number"], "List Number")):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        style_font(st, BODY_SIZE)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(4)
        pf.line_spacing = BODY_LINE
        pf.left_indent = Inches(0.375)
        pf.first_line_indent = Inches(-0.188)
        pPr = st.element.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        numId = OxmlElement("w:numId")
        numId.set(qn("w:val"), str(num_id))
        numPr.append(ilvl)
        numPr.append(numId)
        # 关键：numPr 必须排在 spacing / ind 之前，否则 Word 忽略编号。
        insert_ordered(pPr, numPr, PPR_ORDER)
    return "List Tighter Bullet", "List Tighter Number"


def bullets(doc, style_name, items):
    for it in items:
        p = doc.add_paragraph(style=style_name)
        if isinstance(it, str):
            set_font(p.add_run(clean(it)), size=BODY_SIZE)
        else:
            for text, bold, color in it:
                set_font(p.add_run(clean(text)), size=BODY_SIZE, bold=bold, color=color)


# ----------------------------------------------------------------- tables
def make_table(doc, headers, rows, weights, align=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    # borders
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "C6D3E2")
        borders.append(e)
    tblPr.append(borders)

    def fill_cell(cell, text, bold, color, size, halign, header=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = halign
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.15
        set_font(p.add_run(clean(text)), size=size, bold=bold, color=color)
        if header:
            shade(cell._tc.get_or_add_tcPr(), TABLE_FILL)
        tcPr = cell._tc.get_or_add_tcPr()
        va = OxmlElement("w:vAlign")
        va.set(qn("w:val"), "center")
        tcPr.append(va)

    if align is None:
        align = [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)

    for i, h in enumerate(headers):
        fill_cell(t.rows[0].cells[i], h, True, DARK_BLUE, TABLE_TEXT_SIZE,
                  WD_ALIGN_PARAGRAPH.CENTER, header=True)
    for r, row in enumerate(rows, start=1):
        for i, val in enumerate(row):
            fill_cell(t.rows[r].cells[i], val, False, BLACK, TABLE_TEXT_SIZE, align[i])

    # 跨页行为：表头行在续页重复；任何一行都不允许被拦腰截断；
    # 表头行带 keep-with-next，避免出现「表头孤零零留在页底」。
    for ri, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        insert_ordered(trPr, cant, TRPR_ORDER)
        if ri == 0:
            th = OxmlElement("w:tblHeader")
            insert_ordered(trPr, th, TRPR_ORDER)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True

    widths = column_widths_from_weights(weights, TABLE_WIDTH)
    apply_table_geometry(t, widths, table_width_dxa=TABLE_WIDTH,
                         indent_dxa=TABLE_INDENT, cell_margins_dxa=TABLE_MARGINS)
    return t


def table_note(doc, text):
    return para(doc, text, size=9, color=MUTED, italic=True, before=4, after=12, line=1.15)


# ----------------------------------------------------------------- document
def build(path):
    doc = Document()

    # base styles
    normal = doc.styles["Normal"]
    style_font(normal, BODY_SIZE)
    normal.paragraph_format.space_after = Pt(BODY_AFTER)
    normal.paragraph_format.line_spacing = BODY_LINE

    for name, tok in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        st = doc.styles[name]
        style_font(st, tok["size"], color=RGBColor.from_string(tok["color"]), bold=True)
        st.paragraph_format.space_before = Pt(tok["before"])
        st.paragraph_format.space_after = Pt(tok["after"])
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.keep_with_next = True

    num_ids = add_numbering(doc)
    bullet_style, number_style = add_list_styles(doc, num_ids)

    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    # running header / footer (from page 2)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(hp.add_run(RUNNING_LABEL), size=9, color=MUTED)
    hp.add_run("\t\t")
    set_font(hp.add_run(RUNNING_RIGHT), size=9, color=MUTED)
    hp.paragraph_format.space_after = Pt(0)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run("第 "), size=9, color=MUTED)
    r = fp.add_run()
    set_font(r, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    fr = OxmlElement("w:r")
    frpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    frpr.append(sz)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "595959")
    frpr.append(col)
    fr.append(frpr)
    ft = OxmlElement("w:t")
    ft.text = "1"
    fr.append(ft)
    fld.append(fr)
    fp._p.append(fld)
    set_font(fp.add_run(" 页"), size=9, color=MUTED)
    fp.paragraph_format.space_before = Pt(0)

    # ---------------- masthead ----------------
    para(doc, "会话技术总结", size=10, color=GOLD, bold=True, before=0, after=4)
    para(doc, TITLE_TEXT, size=23, color=INK, bold=True, before=0, after=4, line=1.05)
    para(doc, SUBTITLE_TEXT, size=13, color=RGBColor(0x37, 0x41, 0x51), before=0, after=12, line=1.15)

    meta = [
        ("项目", "智行护航 —— 面向果园小型农机的端侧多模态融合与预测式主动安全头盔系统"),
        ("记录范围", "2026-09-30 ~ 2026-10-01 的开发对话"),
        ("技术栈", "ESP-IDF v6.0.2 / FreeRTOS / C"),
        ("硬件平台", "ESP32-S3-N16R8 ×2（相机板 + 最小板）"),
        ("当前状态", "Phase 1-2 骨架已编译通过；传感器驱动待接入"),
    ]
    for k, v in meta:
        rich(doc, [(k + "：", True, BLACK), (v, False, BLACK)],
             size=10.5, before=0, after=2, line=1.2)
    rule = para(doc, "", before=10, after=2)
    para_bottom_rule(rule)

    # ---------------- 1 ----------------
    doc.add_heading("1. 项目与硬件基线", level=1)
    para(doc, "本项目面向果园等农业生产环境中的小型农机驾驶安全问题，通过头盔上的端侧设备采集多模态"
              "环境信息，在本地完成感知、融合与主动预警。当前手上是两块同型号主控板：一块是集成了 DVP "
              "摄像头底座和 microSD 的相机板（Freenove ESP32-S3 WROOM），另一块是引脚更宽裕的最小系统板。")

    make_table(doc,
        ["角色", "型号", "接口", "关键参数"],
        [
            ["主控（相机板）", "ESP32-S3-WROOM-1 N16R8", "—", "16 MB Flash / 8 MB Octal PSRAM，双核 LX7 @240MHz"],
            ["主控（最小板）", "ESP32-S3-N16R8", "—", "同上，可用 GPIO 明显更多"],
            ["视觉", "OV2640", "DVP + SCCB", "200 万像素；实际可用分辨率需按 PSRAM 带宽核实"],
            ["毫米波雷达", "觅感 MS60-1211S80M", "UART + OUT", "59-64 GHz FMCW，1T2R，探测 ≤10 m，视角 ±60°，3.0-5.5 V / 80 mA"],
            ["惯性传感器", "HW-991", "I2C 或 SPI", "6 轴 16 位；模块供电 5 V 或 3.3 V"],
            ["存储", "板载 microSD", "SDMMC 1-bit", "CLK 39 / CMD 38 / D0 40"],
        ],
        weights=[1.15, 1.75, 1.05, 2.55])
    table_note(doc, "表 1  硬件基线。雷达与 IMU 的型号均以技术资料中的产品说明书为准。")

    # ---------------- 2 ----------------
    doc.add_heading("2. 关键技术结论", level=1)

    doc.add_heading("2.1 相机板的引脚预算已经见底", level=2)
    para(doc, "把摄像头、SD 卡、PSRAM、USB、串口和板载彩灯都扣掉之后，相机板只剩 5 个可用 GPIO，"
              "其中 GPIO2 还被板载 LED 占用。这个数字直接决定了后面单板还是双板的架构选择。")
    make_table(doc,
        ["用途", "引脚", "说明"],
        [
            ["摄像头（DVP）", "4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 15, 16, 17, 18", "对应 Arduino 例程里的 CAMERA_MODEL_ESP32S3_EYE"],
            ["microSD", "38, 39, 40", "SDMMC 1-bit 模式，官方注释标注不可修改"],
            ["Octal PSRAM", "35, 36, 37", "ESP32-S3-WROOM-1 数据手册明确：已接 OSPI PSRAM，不可作他用"],
            ["USB / 串口日志 / 彩灯", "19, 20 / 43, 44 / 48", "USB_D±、U0TXD-RXD、WS2812"],
            ["剩余可用", "1, 2, 14, 21, 47", "GPIO2 已被板载 LED 占用，实际可外接的只有 1 / 14 / 21 / 47"],
        ],
        weights=[1.4, 2.5, 2.6])
    table_note(doc, "表 2  相机板引脚占用。资料包里 CameraWebServer 示例的 board_config.h 选的是 "
                    "CAMERA_MODEL_ESP_EYE，那是 ESP32 的引脚表，用在 S3 上是错的。")

    doc.add_heading("2.2 IMU 的核心其实是 Bosch BMI270", level=2)
    para(doc, "HW-991 产品说明书只教了「装 Arduino 的 BMI270 库跑 Basic 例子」，没有给寄存器手册。"
              "从引脚命名（ADO/MISO、SDA/MOSI、SCL/SCLK、CS、INT1/INT2，以及 ASDx/ASCx 辅助接口）可以确认，"
              "板载芯片是 Bosch BMI270 六轴 IMU，支持 I2C 与 SPI 二选一。"
              "在 ESP-IDF 下需要 Bosch 官方 BMI270 SensorAPI（C 版本），而不是 Arduino 库。")

    doc.add_heading("2.3 雷达协议要点", level=2)
    para(doc, "雷达的核心是隔空科技的 AT6010 SoC，资料包里那份《AT6010 SOC HCI Protocol V1.0》"
              "把帧格式写得很清楚。上位机命令帧以 0x58 开头，雷达回复帧以 0x59 开头，"
              "两者结构相同：命令字 + 参数长度 + 参数 + 校验和 + 0x00 结尾，"
              "校验和等于前面所有字节之和（取低 8 位）。默认波特率是 921600 bps，可切到 9600 或 115200。")
    para(doc, "真正要在驱动里解析的是主动上报帧：帧头固定 0x5A，后面跟「长度 + 载荷」，"
              "载荷的第一字节是类型码，校验和覆盖帧头、长度和载荷三部分。"
              "关键行为是：只有检测到目标才会输出，没有目标时完全不发数据，"
              "所以驱动必须靠超时来判断「目标消失」，不能把「收不到」当成「没目标」。")
    make_table(doc,
        ["TYPE", "内容", "载荷长度", "主要字段"],
        [
            ["0", "完整检测信息", "21 字节", "is_detected、det_result、距离(mm)、角度、速度(预留)、距离/角度置信度、帧号"],
            ["1", "测高", "5 字节", "身高(mm)、状态位（校准完成 / 进入区域 / 测量完成）"],
            ["2", "占位检测", "5 字节", "RB 编号、原始功率、距离补偿后功率、超阈值邻域数"],
            ["3", "运动存在", "9 字节", "is_detected、det_result、距离、角度、速度(预留)"],
            ["4", "呼吸心率", "9 字节", "检测结果、呼吸率、心率、角度(保留)、距离"],
            ["5", "分区检测", "17 字节", "目标数 + 最多 3 个分区内最近目标的距离与角度"],
        ],
        weights=[0.6, 1.5, 0.9, 3.5],
        align=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
               WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    table_note(doc, "表 3  雷达主动上报类型。det_result 是位域：0x01 靠近 / 0x02 远离 / 0x04 运动 / "
                    "0x08 微动 / 0x10 呼吸。")

    doc.add_heading("2.4 三条必须先知道的限制", level=2)

    doc.add_heading("雷达不输出速度", level=3)
    para(doc, "产品手册宣传「输出目标编号、距离、速度、角度」，但 HCI 协议里 velo_val 字段在 TYPE 0、"
              "TYPE 3 两处结构体定义中都标注「目前预留 / todo」。也就是说速度拿不到。"
              "原计划里「距离过近 + 高速接近 = 危险」这条规则，必须改成用 Δ距离/Δt 自己估算趋势，"
              "算法设计要提前调整。")

    doc.add_heading("它是人体存在感应雷达，不是防撞雷达", level=3)
    para(doc, "这款雷达的设计场景是吸顶或贴墙安装（盲区 0.1 m、探测 ≤10 m），输出的是单一目标的距离和角度，"
              "没有目标编号，也没有多目标跟踪——「多目标定位」只体现在 TYPE 5 的「每个分区输出最近目标，最多 3 个」。"
              "产品定位说成「接近预警」比「碰撞预测」更符合它的实际能力。")

    doc.add_heading("ESP-IDF v6.0.2 删掉了旧驱动 API", level=3)
    para(doc, "实测确认：这套 IDF 里 driver/i2c.h 和 driver/rmt.h 已经不存在，只剩 driver/i2c_master.h "
              "和 driver/rmt_tx.h。网上和资料包里绝大多数例程（包括 Freenove 那批 Arduino 例程）"
              "在 v6 下直接编译不过，写代码必须按 v6 的新 API。")

    # ---------------- 3 ----------------
    doc.add_heading("3. 技术路线决策", level=1)

    doc.add_heading("3.1 双板分工：按传感器归属划分，而不是「采集 vs 处理」", level=2)
    para(doc, "最初的想法是「一块板采集数据、一块板处理数据」，但图像没法低成本搬过去："
              "QVGA 一帧 RGB565 就是 320×240×2 ≈ 150 KB，10 fps 相当于 1.5 MB/s，"
              "而 UART 在 921600 bps 下只有约 92 KB/s，连 1 fps 都跑不动；"
              "走 WiFi 虽然能传，但 ESP32-S3 上「摄像头 + 推理 + WiFi」同时跑，带宽和实时性都不可控。")
    para(doc, "所以现实的分工是：视觉模型必须留在相机板上跑，另一块板只收结构化的检测结果（几十字节），"
              "融合决策集中在一个节点上。")
    make_table(doc,
        ["维度", "相机板（视觉节点）", "最小板（主控 / 融合节点）"],
        [
            ["传感器", "OV2640 + microSD", "雷达 UART + IMU I2C"],
            ["模型", "视觉模型（分类 / 检测）", "轻量时序模型；建议先做阈值规则"],
            ["对外输出", "结构化结果：帧号、时间戳、类别、置信度、方位", "融合结论 + 报警动作"],
            ["角色", "只上报观测，不做决策", "唯一决策点，出危险就报警"],
        ],
        weights=[0.95, 2.6, 2.95])
    table_note(doc, "表 4  双板分工。板间只需要传几十字节的检测结果，UART 完全够用。")

    doc.add_heading("3.2 ESP32-S3 的算力现实", level=2)
    para(doc, "ESP32-S3 是双核 LX7 @240MHz，有向量指令但没有 NPU，所有推理都吃 CPU。"
              "下表的数字是经验量级，用来判断方案可行性，实际必须自己实测。")
    make_table(doc,
        ["模型层级", "输入尺寸", "单帧耗时", "可用帧率"],
        [
            ["分类（MobileNet 级）", "96-112", "30-80 ms", "12-30 fps"],
            ["检测（YOLO / ESPDet 级）", "160-192", "200-800 ms", "1-5 fps"],
            ["人脸检测级", "QVGA", "100-300 ms", "3-10 fps"],
            ["IMU 时序分类（1D-CNN）", "6 通道 × 50 点", "< 5 ms", "不是瓶颈"],
        ],
        weights=[2.0, 1.4, 1.3, 1.8],
        align=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
               WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    table_note(doc, "表 5  预期性能量级。检测级模型能跑，但只能低分辨率低帧率，"
                    "对「前方有人/有车接近」这类持续数百毫秒到数秒的预警够用。")

    doc.add_heading("3.3 建议的推进顺序", level=2)
    bullets(doc, number_style, [
        "先在相机板上把摄像头跑起来，实测视觉模型的真实帧率——这是整个方案最大的未知数，先把它变成已知数。",
        "同时用最小板做 Phase 3（雷达 + IMU 接入），先出数据。",
        "两块板各自独立跑通后，再定义板间协议（UART，字节帧 + 时间戳 + 序号 + 校验）。",
        "最后再决定要不要合并回单板。",
    ])

    # ---------------- 4 ----------------
    doc.add_heading("4. 固件工程现状", level=1)
    para(doc, "已在 G:\\codex-workspace\\firmware 建好 ESP-IDF 工程骨架，落实 PLAN.md 的 Phase 1 和 Phase 2："
              "LED 心跳指示 + Task → Queue → Task 数据通路，并让 Queue、Semaphore、Mutex 三种同步原语"
              "都有真实落点。传感器数据目前由模拟任务产生，Phase 3 接入真实驱动时只需替换这一个任务，"
              "融合层代码不用改。")
    make_table(doc,
        ["任务", "职责", "优先级", "栈（字）", "通信方式"],
        [
            ["radar_sim", "模拟雷达周期采集（Phase 3 替换为真实驱动）", "5", "3072", "→ sensor_queue"],
            ["fusion", "融合 + 风险判定 + 统计", "4", "4096", "收 sensor_queue；Mutex 保护统计；xQueueOverwrite 写 led_queue"],
            ["button", "响应 BOOT 按键事件", "3", "3072", "等 button_sem（ISR 给出）"],
            ["led", "按风险等级驱动板载 LED", "2", "2048", "收 led_queue（深度 1，只取最新值）"],
        ],
        weights=[0.9, 2.2, 0.7, 0.8, 1.9],
        align=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
               WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    table_note(doc, "表 6  固件任务设计。用 ESP-IDF v6.0.2 实测编译通过，产物 zhixing_hud.bin 约 183 KB，"
                    "4 MB 的 app 分区用了 4%，零警告。")

    para(doc, "引脚全部集中在 board_config.h，换板或改接线只改这一个文件。传感器预留接线为："
              "雷达走 UART1（TX=GPIO14，RX=GPIO21，默认 921600），IMU 走 I2C0（SDA=GPIO1，SCL=GPIO47）。"
              "这两组线目前尚未实际连接。")

    # ---------------- 5 ----------------
    doc.add_heading("5. 环境问题与修复记录", level=1)
    para(doc, "在搭工程的过程中顺带做了一次开发环境体检，发现并处理了下面这些问题。"
              "其中前两项会直接导致「找不到 IDF」「找不到编译工具」，属于必须先解决的。")
    make_table(doc,
        ["问题", "原因", "处理", "状态"],
        [
            ["ESP-IDF 扩展报找不到 IDF", "settings.json 里 espIdfPathWin 里的 v 被写成转义序列 \\u000b（垂直制表符），解析成 D:\\esp + 控制字符 + 6.0.2\\esp-idf，路径不存在",
             "定点替换为 D:\\esp\\v6.0.2\\esp-idf，JSON 校验通过", "已修复"],
            ["IDF 编译找不到 ninja/ccache", "Agent 沙箱会重置 python 子进程的 PATH，idf.py 启动的 cmake 拿不到工具链路径",
             "改为在沙箱外执行编译；你自己的终端不受影响", "已规避"],
            ["OpenOCD 板级配置对不上", "openOcdConfigs 指向 board/esp32-wrover-kit-3.3v.cfg，那是 ESP32 的板",
             "改为 board/esp32s3-builtin.cfg", "已修复"],
            ["控制台中文乱码", "输出编码 UTF-8、输入编码 GB2312、活动代码页 936 三者不一致；且系统未装 PowerShell 7，终端落到 5.1（默认 $OutputEncoding 为 ASCII）",
             "安装 PowerShell 7.6.6；写入 PS7/PS5.1 两个 profile 统一 UTF-8；设置 PYTHONUTF8=1", "已修复"],
            ["深目录构建可能失败", "LongPathsEnabled=0，git core.longpaths 未设",
             "git core.longpaths 已设 true；系统级 LongPathsEnabled 需要管理员，尚未执行", "部分完成"],
            ["C 盘只剩 9.53 GB", "休眠文件 12.49 GB + 系统更新残留 + 各类缓存",
             "两轮白名单清理，见第 6 节", "已处理"],
        ],
        weights=[1.35, 2.75, 2.05, 0.85],
        align=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
               WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
    table_note(doc, "表 7  环境问题与修复记录。")

    # ---------------- 6 ----------------
    doc.add_heading("6. C 盘空间清理记录", level=1)
    para(doc, "清理遵循「只删程序会自动重建的缓存和系统已废弃的临时目录」这一条原则，"
              "所有目标写死在脚本白名单里，不用通配符，删除前先备份、删除后逐项核对。"
              "C 盘可用空间从 9.53 GB 恢复到 28.94 GB。")
    para(doc, "实际删除的包括：休眠文件 hiberfil.sys（12.49 GB，通过 powercfg /h off）、"
              "Windows 更新残留 $WINDOWS.~BT 与 $WinREAgent（合计 1.47 GB）、"
              "ESP-IDF 安装器的离线包缓存 eim\\offline_archives（3.09 GB）、"
              "VS Code 的扩展安装包缓存与 C++ 智能感知数据库（合计约 2.6 GB）、"
              "npm 缓存、Edge 与 VS Code 的着色器/代码缓存，以及显卡驱动安装包缓存。")
    para(doc, "明确没有碰的包括：C:\\eSupport（6.12 GB 出厂驱动包）、Arduino15（8.09 GB）、"
              "Windows\\Installer（3.87 GB）、DriverStore（10.61 GB）、.espressif 工具链（6.79 GB）、"
              "页面文件，以及文档、桌面、视频、下载和聊天记录等全部用户数据。")
    callout(doc, "唯一的行为变更：",
            "休眠与快速启动已随 hiberfil.sys 一起关闭。如果需要，可用 powercfg /h /type reduced "
            "只恢复快速启动（约占 6 GB），或用 powercfg /h on 完整恢复（约占 12.5 GB）。")
    para(doc, "所有原始备份、变更清单和还原脚本已归集到 G:\\codex-workspace\\backup，"
              "其中 restore-env.ps1 可以逐项还原配置改动。")

    # ---------------- 7 ----------------
    doc.add_heading("7. 待办与下一步", level=1)
    bullets(doc, bullet_style, [
        [("雷达协议驱动：", True, DARK_BLUE), ("按 0x5A 主动上报帧实现 UART 解析，含超时判定与置信度过滤。", False, BLACK)],
        [("IMU 驱动：", True, DARK_BLUE), ("接入 Bosch BMI270 SensorAPI，先出原始六轴数据，再做姿态解算。", False, BLACK)],
        [("实测视觉帧率：", True, DARK_BLUE), ("这是双板方案能否成立的关键未知数，建议优先验证。", False, BLACK)],
        [("系统级长路径：", True, DARK_BLUE), ("管理员运行 tools/admin-fix.ps1 开启 LongPathsEnabled。", False, BLACK)],
        [("待确认的硬件事实：", True, DARK_BLUE), ("雷达与 IMU 的供电电压（建议统一 3.3 V）、最小板的实际型号与引脚表。", False, BLACK)],
        [("架构文档同步：", True, DARK_BLUE), ("ARCHITECTURE.md 的架构图与正文自相矛盾（Fusion/Decision 职责重叠、漏画 Queue），建议按本次结论重画。", False, BLACK)],
    ])

    doc.add_heading("附录  本次会话产出的文件", level=1)
    bullets(doc, bullet_style, [
        "firmware\\ —— ESP-IDF 工程骨架（已编译通过）",
        "firmware\\main\\board_config.h —— 板型与引脚定义，全部引脚集中在此",
        "tools\\clean-c-drive.ps1 —— C 盘白名单清理脚本（默认预览，-Apply 才执行）",
        "tools\\admin-fix.ps1 —— 需要管理员权限的环境修复脚本（长路径等）",
        "backup\\ —— 环境修复的全部安全备份与还原脚本",
    ])

    doc.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1]
    os.makedirs(os.path.dirname(out), exist_ok=True)
    build(out)
    print("OK ->", out)
